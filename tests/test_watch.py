"""`jjh watch`: incremental cycles over a fake LinkedIn, scheduling, and the loop with a fake clock."""

from __future__ import annotations

import io
import random
from contextlib import redirect_stdout
from datetime import datetime, timedelta
from urllib.parse import parse_qs, urlsplit

import httpx
import pytest
from openpyxl import load_workbook

from jev_job_hunter import linkedin, watch
from jev_job_hunter.profile import load_profile
from jev_job_hunter.watch import WatchState, next_run_at, run_cycle, watch_key

from test_linkedin import CARDS, ROOT, FakeLinkedIn, fake_client, profile  # noqa: F401  (profile is a fixture)

T0 = datetime(2026, 10, 14, 10, 0)
QUERIES = ["Backend Engineer", "AI Engineer"]  # profile.example.yaml search.queries


def card(i: int) -> dict:
    return {**CARDS[0], "id": str(5000 + i), "title": f"Backend Engineer {i}",
            "href": f"https://www.linkedin.com/jobs/view/backend-{5000 + i}"}


@pytest.fixture
def jev_calls(monkeypatch):
    calls = []
    real = linkedin.ask

    def counting(state, questions, mock):
        calls.append(sorted(questions))
        return real(state, questions, mock)

    monkeypatch.setattr(linkedin, "ask", counting)
    return calls


def _cycle(tmp_path, prof, site, at: datetime) -> dict:
    state = WatchState.load(tmp_path, watch_key(QUERIES, "South Korea", "day", "any"))
    with redirect_stdout(io.StringIO()):
        return run_cycle(tmp_path, state, prof, QUERIES, "South Korea", "day", "any",
                         max_open=10, delay=0, mock=True, client=fake_client(site), clock=lambda: at)


def _search_starts(site) -> list[str]:
    return [parse_qs(urlsplit(u).query)["start"][0] for u in site.visited if u.startswith(linkedin.GUEST_SEARCH)]


def _postings_read(site) -> int:
    return sum(1 for u in site.visited if "/jobPosting/" in u)


def test_cycles_score_only_new_jobs(tmp_path, profile, jev_calls):
    prof = load_profile(ROOT, str(profile))

    site = FakeLinkedIn()
    first = _cycle(tmp_path, prof, site, T0)
    assert (first["new"], first["scored"], first["error"]) == (3, 2, "")
    assert jev_calls and _postings_read(site) >= 2

    jev_calls.clear()
    site = FakeLinkedIn()
    second = _cycle(tmp_path, prof, site, T0 + timedelta(minutes=45))
    assert (second["new"], second["scored"]) == (0, 0)
    assert jev_calls == [] and _postings_read(site) == 0
    assert second["requests"] == 2  # one search page per query, nothing else

    site = FakeLinkedIn(cards=[card(1), *CARDS])
    third = _cycle(tmp_path, prof, site, T0 + timedelta(minutes=90))
    assert third["new"] == 1 and (third["listings"], third["matches"]) == (4, 3)

    wb = load_workbook(third["xlsx"])
    ws = wb["All listings"]
    head = [c.value for c in ws[1]]
    rows = {r[head.index("Job ID")].value: r for r in ws.iter_rows(min_row=2)}
    assert len(rows) == 4
    assert rows["5001"][head.index("First seen")].value == "2026-10-14 11:30"
    assert rows["4001"][head.index("First seen")].value == "2026-10-14 10:00"
    assert rows["5001"][head.index("Title")].font.bold and not rows["4001"][head.index("Title")].font.bold

    saved = WatchState.load(tmp_path, watch_key(QUERIES, "South Korea", "day", "any"))
    assert set(saved.seen) == {"4001", "4002", "4003", "5001"} and len(saved.cycles) == 3


def test_paging_goes_deeper_only_while_everything_is_new(tmp_path, profile):
    prof = load_profile(ROOT, str(profile))
    site = FakeLinkedIn(cards=[card(i) for i in range(35)])
    with redirect_stdout(io.StringIO()):
        _cycle(tmp_path, prof, site, T0)
    # query 1: 3 full pages of new jobs, then the 30-card cap; query 2 sees the same jobs on page 1 → stop
    assert _search_starts(site) == ["0", "10", "20", "0"]

    site = FakeLinkedIn(cards=[card(100), card(101)] + [card(i) for i in range(35)])
    with redirect_stdout(io.StringIO()):
        got = _cycle(tmp_path, prof, site, T0 + timedelta(hours=1))
    assert _search_starts(site) == ["0", "0"] and got["new"] == 2


class BlockedPostings(FakeLinkedIn):
    def __call__(self, request):
        if "/jobPosting/" in str(request.url):
            self.visited.append(str(request.url))
            return httpx.Response(429)
        return super().__call__(request)


def test_blocked_cycle_keeps_progress_and_retries_unread(tmp_path, profile, monkeypatch):
    monkeypatch.setattr(linkedin, "RATE_LIMIT_BACKOFF", 0.0)
    prof = load_profile(ROOT, str(profile))
    blocked = _cycle(tmp_path, prof, BlockedPostings(), T0)
    assert "rate limit" in blocked["error"] and blocked["scored"] == 0
    assert blocked["new"] == 3 and blocked["pending"] == 2  # the two picked-but-unread jobs wait
    assert load_workbook(blocked["xlsx"])["All listings"].max_row == 4

    site = FakeLinkedIn()
    retry = _cycle(tmp_path, prof, site, T0 + timedelta(hours=1))
    assert (retry["new"], retry["scored"], retry["pending"], retry["error"]) == (0, 2, 0, "")
    assert _postings_read(site) >= 2  # read straight from the queue, no re-triage needed


def _cycle_n(tmp_path, prof, site, at, max_open) -> dict:
    state = WatchState.load(tmp_path, watch_key(QUERIES, "South Korea", "day", "any"))
    with redirect_stdout(io.StringIO()):
        return run_cycle(tmp_path, state, prof, QUERIES, "South Korea", "day", "any",
                         max_open=max_open, delay=0, mock=True, client=fake_client(site), clock=lambda: at)


def test_reads_are_spread_across_cycles(tmp_path, profile, jev_calls):
    prof = load_profile(ROOT, str(profile))
    first = _cycle_n(tmp_path, prof, FakeLinkedIn(), T0, max_open=1)
    assert (first["new"], first["scored"], first["pending"]) == (3, 1, 1)

    jev_calls.clear()
    site = FakeLinkedIn()
    second = _cycle_n(tmp_path, prof, site, T0 + timedelta(minutes=40), max_open=1)
    assert (second["new"], second["scored"], second["pending"]) == (0, 1, 0)
    assert _postings_read(site) >= 1
    assert all("open_0" not in q for q in jev_calls)  # the queued job was not triaged again
    state = WatchState.load(tmp_path, watch_key(QUERIES, "South Korea", "day", "any"))
    assert sum(1 for c in state.listings if c.get("opened")) == 2 and len(state.matches) == 2


def test_state_file_per_search(tmp_path):
    a = WatchState.load(tmp_path, "aaaa1111")
    a.seen["1"] = "x"
    a.save()
    assert a.path.name == "linkedin-watch.json"
    assert WatchState.load(tmp_path, "aaaa1111").seen == {"1": "x"}
    with redirect_stdout(io.StringIO()):
        b = WatchState.load(tmp_path, "bbbb2222")
    assert b.path.name == "linkedin-watch-bbbb2222.json" and b.seen == {}
    assert b.xlsx == tmp_path / "results" / "linkedin-watch-bbbb2222.xlsx"


def test_next_run_at_gaps_quiet_hours_and_backoff():
    rng = random.Random(7)
    gaps = [(next_run_at(T0, (30, 60), None, 1, rng) - T0).total_seconds() / 60 for _ in range(200)]
    assert min(gaps) >= 30 and max(gaps) <= 60 and len({round(g) for g in gaps}) > 20

    late = datetime(2026, 10, 14, 0, 45)  # +30–60 min lands in 01–07 → pushed to 07:00–07:20
    for _ in range(50):
        t = next_run_at(late, (30, 60), (1, 7), 1, rng)
        assert t.date() == late.date() and t.hour == 7 and t.minute <= 20

    assert all(60 <= (next_run_at(T0, (30, 60), None, 2, rng) - T0).total_seconds() / 60 <= 120 for _ in range(50))
    assert all((next_run_at(T0, (30, 60), None, 8, rng) - T0).total_seconds() / 60 == 240 for _ in range(20))

    night = datetime(2026, 10, 14, 23, 30)  # quiet 22–6 wraps midnight → next morning 06:xx
    t = next_run_at(night, (30, 60), (22, 6), 1, rng)
    assert t.date() == night.date() + timedelta(days=1) and t.hour == 6


def test_watch_loop_waits_runs_and_reschedules(tmp_path, profile, monkeypatch):
    (tmp_path / "config").mkdir()
    (tmp_path / "config" / "profile.yaml").write_text(profile.read_text(encoding="utf-8"), encoding="utf-8")
    now = [T0]
    slept = []

    def sleeper(s):
        slept.append(s)
        now[0] += timedelta(seconds=s)

    buf = io.StringIO()
    with redirect_stdout(buf):
        watch.watch(tmp_path, location="South Korea", max_open=10, delay=0, mock=True, every=(30, 60),
                    quiet=None, client_factory=fake_client, clock=lambda: now[0], sleeper=sleeper, max_cycles=2)
    out = buf.getvalue()
    assert out.count("WATCH  2026-10-14") == 2 and "Next" in out and "watch stopped" in out
    waited = sum(slept) / 60
    assert 30 <= waited <= 60 and max(slept) <= watch.TICK  # one gap, slept in short slices
    state = WatchState.load(tmp_path, watch_key(QUERIES, "South Korea", "day", "any"))
    assert [c["new"] for c in state.cycles] == [3, 0]


def test_once_skips_quiet_hours(tmp_path, profile):
    site = FakeLinkedIn()
    buf = io.StringIO()
    with redirect_stdout(buf):
        watch.watch(ROOT, profile_path=str(profile), max_open=10, delay=0, mock=True, quiet=(1, 7), once=True,
                    client_factory=lambda: fake_client(site), clock=lambda: datetime(2026, 10, 14, 3, 0))
    assert "quiet hours" in buf.getvalue() and site.visited == []


def test_each_find_is_saved_as_it_arrives_and_a_crash_keeps_them(tmp_path, profile, monkeypatch):
    prof = load_profile(ROOT, str(profile))
    key = watch_key(QUERIES, "South Korea", "day", "any")
    on_disk = []

    def boom(*a, **k):  # Jev fails after the search: the cards must already be on disk
        on_disk.append(len(WatchState.load(tmp_path, key).listings))
        raise RuntimeError("Jev unreachable")

    monkeypatch.setattr(watch, "triage", boom)
    got = _cycle(tmp_path, prof, FakeLinkedIn(), T0)
    assert on_disk == [3] and got["new"] == 3 and "Jev unreachable" in got["error"]
    saved = WatchState.load(tmp_path, key)
    assert len(saved.listings) == 3 and set(saved.seen) == {"4001", "4002", "4003"}
    assert all(c["query"] and c["first_seen"] == "2026-10-14 10:00" for c in saved.listings)


def test_dismissed_jobs_are_not_read_from_the_queue(tmp_path, profile):
    prof = load_profile(ROOT, str(profile))
    first = _cycle_n(tmp_path, prof, FakeLinkedIn(), T0, max_open=1)
    assert first["pending"] == 1
    waiting = WatchState.load(tmp_path, watch_key(QUERIES, "South Korea", "day", "any")).pending[0]["id"]
    watch.set_dismissed(tmp_path, [waiting])
    site = FakeLinkedIn()
    second = _cycle_n(tmp_path, prof, site, T0 + timedelta(minutes=40), max_open=1)
    assert (second["scored"], second["pending"]) == (0, 0) and _postings_read(site) == 0


def test_jobs_found_before_logos_get_one_when_seen_again(tmp_path, profile):
    prof = load_profile(ROOT, str(profile))
    _cycle(tmp_path, prof, FakeLinkedIn(), T0)  # cards without logos, like jobs saved before logos were kept
    logo = "https://media.licdn.com/dms/image/v2/x/company-logo_100_100/toss_logo?e=1&v=beta"
    site = FakeLinkedIn(cards=[{**CARDS[0], "logo": logo}, *CARDS[1:]])
    again = _cycle(tmp_path, prof, site, T0 + timedelta(minutes=45))
    assert again["new"] == 0 and _postings_read(site) == 0  # nothing extra fetched for it
    saved = {c["id"]: c for c in WatchState.load(tmp_path, watch_key(QUERIES, "South Korea", "day", "any")).listings}
    assert saved["4001"]["logo"] == logo and not saved["4002"].get("logo")


def test_reposts_under_a_new_job_id_are_not_new(tmp_path, profile):
    prof = load_profile(ROOT, str(profile))
    repost = {**CARDS[0], "id": "4998", "title": "Backend Engineer!", "posted_text": "1 hour ago"}
    first = _cycle(tmp_path, prof, FakeLinkedIn(cards=[*CARDS, repost]), T0)  # same search, same cycle
    assert first["new"] == len(CARDS)
    later = {**CARDS[0], "id": "4999", "href": "https://www.linkedin.com/jobs/view/backend-engineer-4999"}
    again = _cycle(tmp_path, prof, FakeLinkedIn(cards=[later, *CARDS[1:]]), T0 + timedelta(minutes=45))
    assert again["new"] == 0


def test_reposts_saved_earlier_are_dropped_on_load(tmp_path):
    dup = {**CARDS[0], "id": "4998"}
    st = WatchState(tmp_path / "w.json", "k", {"listings": [CARDS[0], dup, {**CARDS[0], "id": "4997"}],
                                                "matches": [{"id": "4997", "fit": 0.5}],
                                                "pending": [dup, CARDS[0]]})
    assert [c["id"] for c in st.listings] == ["4001", "4997"]  # the scored one stays
    assert [c["id"] for c in st.pending] == ["4001"]

