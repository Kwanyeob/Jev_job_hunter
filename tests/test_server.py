"""`jjh serve` web app: invite login, per-user isolation, resume, profile, watch controls, export."""

from __future__ import annotations

import io
import json
import random
from contextlib import redirect_stdout
from datetime import datetime, timedelta

import pytest
from starlette.testclient import TestClient

from jev_job_hunter.resume import ResumeError
from jev_job_hunter.scheduler import Hub, Scheduler
from jev_job_hunter.server import COOKIE, create_app
from jev_job_hunter.users import Users

from test_linkedin import FakeLinkedIn, fake_client

PROFILE = {
    "candidate": {"headline": "Backend Engineer", "summary": "APIs.", "years_experience": 4, "seniority": "mid",
                  "skills": ["Python", "Go"], "target_roles": ["Backend Engineer", "AI Engineer"],
                  "interests": [], "locations": ["Seoul"], "languages": ["Korean"], "education": "",
                  "email": "leak@example.com"},
    "search": {"queries": ["Backend Engineer", "AI Engineer", "Platform Engineer", "Data Engineer"],
               "location": "Seoul, South Korea", "remote": "any", "posted": "week"},
}


@pytest.fixture
def svc(tmp_path, monkeypatch):
    monkeypatch.delenv("SECRET_KEY", raising=False)
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    users = Users(tmp_path)
    hub = Hub()
    sched = Scheduler(users, hub, mock=True, clock=lambda: datetime(2026, 10, 16, 10, 0), quiet=None,
                      client_factory=lambda: fake_client(FakeLinkedIn()), delay=0, pause=lambda s: None,
                      rng=random.Random(1))
    calls = {"extract": 0, "reader": True}

    def extract(pdf: bytes) -> dict:
        calls["extract"] += 1
        if not pdf.startswith(b"%PDF"):
            raise ResumeError("not a PDF file")
        return json.loads(json.dumps(PROFILE))

    app = create_app(tmp_path, scheduler=sched, hub=hub, extract=extract, reader_ready=lambda: calls["reader"])
    return app, users, sched, tmp_path, calls


def login(app, users: Users, name: str) -> TestClient:
    _, token = users.invite(name)
    client = TestClient(app)
    r = client.get(f"/invite/{token}", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/" and COOKIE in r.cookies
    return client


def test_needs_an_invite(svc):
    app, users, *_ = svc
    anon = TestClient(app)
    assert anon.get("/").status_code == 401
    r = anon.get("/api/state")
    assert r.status_code == 401 and "invite" in r.json()["error"]  # no language hint: English
    assert "초대" in anon.get("/api/state", headers={"Accept-Language": "ko-KR,ko"}).json()["error"]
    anon.cookies.set("jjh_lang", "ko")  # the page's language switch wins over the browser
    assert "초대" in anon.get("/api/state", headers={"Accept-Language": "en"}).json()["error"]
    assert anon.get("/invite/not-a-token", follow_redirects=False).status_code == 403
    anon.cookies.set(COOKIE, "forged.1.2.3")
    assert anon.get("/api/state").status_code == 401


def test_invite_login_then_app(svc):
    app, users, *_ = svc
    client = login(app, users, "Minsu")
    page = client.get("/")
    assert page.status_code == 200 and "Shortlist" in page.text and "자동 수집" in page.text
    st = client.get("/api/state").json()
    assert st["user"] == "Minsu" and st["profile"] is None and st["watch"]["enabled"] is False
    assert st["jev"] is False and st["max_queries"] == 3


def test_resume_fills_profile_without_storing_pdf_or_contacts(svc):
    app, users, sched, data_dir, calls = svc
    client = login(app, users, "Minsu")
    r = client.post("/api/resume", content=b"%PDF-1.7 fake resume", headers={"Content-Type": "application/pdf"})
    assert r.status_code == 200
    prof = r.json()["profile"]
    assert prof["search"]["queries"] == ["Backend Engineer", "AI Engineer", "Platform Engineer"]  # capped at 3
    assert "email" not in prof["candidate"]
    assert client.get("/api/state").json()["profile"]["candidate"]["headline"] == "Backend Engineer"
    assert not list(data_dir.rglob("*.pdf"))
    saved = next(data_dir.rglob("profile.yaml")).read_text(encoding="utf-8")
    assert "leak@example.com" not in saved

    bad = client.post("/api/resume", content=b"hello", headers={"Content-Type": "application/pdf"})
    assert bad.status_code == 400 and "resume" in bad.json()["error"]
    calls["reader"] = False
    off = client.post("/api/resume", content=b"%PDF-1.7", headers={"Content-Type": "application/pdf"})
    assert off.status_code == 503 and calls["extract"] == 2


def test_profile_save_validates(svc):
    app, users, *_ = svc
    client = login(app, users, "Minsu")
    empty = {"candidate": {"target_roles": []}, "search": {}}
    assert client.post("/api/profile", json=empty).status_code == 400
    assert client.post("/api/profile", content=b"{nope", headers={"Content-Type": "application/json"}).status_code == 400
    ok = client.post("/api/profile", json=PROFILE)
    assert ok.status_code == 200 and len(ok.json()["profile"]["search"]["queries"]) == 3


def test_watch_controls_and_isolation(svc):
    app, users, sched, data_dir, _ = svc
    a, b = login(app, users, "A"), login(app, users, "B")
    assert a.post("/api/watch", json={"enabled": True}).status_code == 400  # no profile yet
    assert a.post("/api/run-now").status_code == 400
    a.post("/api/profile", json=PROFILE)
    on = a.post("/api/watch", json={"enabled": True}).json()
    assert on["enabled"] and on["next_run"]
    assert b.get("/api/state").json()["watch"]["enabled"] is False

    assert a.get("/api/export.xlsx").status_code == 404
    uid_a = users.find("A").id
    with redirect_stdout(io.StringIO()):
        sched.run_one(uid_a)  # watch on: the search only lists and triages; reading waits for a read run
    res = a.get("/api/results").json()
    assert res["total"] == 3 and res["scored"] == 0 and a.get("/api/state").json()["watch"]["pending"] == 2
    with redirect_stdout(io.StringIO()):
        sched.run_one(uid_a, "read")
    res = a.get("/api/results").json()
    assert res["total"] == 3 and res["scored"] == 2 and res["last_at"] == "2026-10-16 10:00"
    assert {r["first_seen"] for r in res["rows"]} == {"2026-10-16 10:00"}
    x = a.get("/api/export.xlsx")
    assert x.status_code == 200 and "attachment" in x.headers["content-disposition"]
    # B sees none of it
    assert b.get("/api/results").json() == {"rows": [], "total": 0, "scored": 0, "last_at": None}
    assert b.get("/api/export.xlsx").status_code == 404
    assert b.get("/api/state").json()["profile"] is None

    off = a.post("/api/watch", json={"enabled": False}).json()
    assert off["enabled"] is False and off["next_run"] is None
    # "search now" rests 10 minutes after a search, so searches never bunch up
    soon = a.post("/api/run-now")
    assert soon.status_code == 429 and "10:10" in soon.json()["error"]
    assert a.get("/api/state").json()["watch"]["run_now_at"] == "2026-10-16T10:10:00"
    sched.clock = lambda: datetime(2026, 10, 16, 10, 10)
    assert a.get("/api/state").json()["watch"]["run_now_at"] is None
    assert a.post("/api/run-now").status_code == 200
    assert a.get("/api/state").json()["watch"]["queued"] is True


def test_revoke_ends_session(svc):
    app, users, *_ = svc
    client = login(app, users, "Minsu")
    assert client.get("/api/state").status_code == 200
    users.revoke("Minsu")
    assert client.get("/api/state").status_code == 401


def test_hub_since_and_rollover():
    hub = Hub(keep=3)
    for i in range(5):
        hub.emit("u1", "status", {"i": i})
    hub.emit("u2", "status", {"i": 99})
    got = hub.since("u1", 0)
    assert [e["id"] for e in got] == [3, 4, 5] and hub.since("u1", 4)[0]["data"] == {"i": 4}
    assert [e["data"]["i"] for e in hub.since("u2", 0)] == [99]


def test_dismissed_jobs_leave_results_and_excel_and_can_come_back(svc):
    app, users, sched, _, _ = svc
    a, b = login(app, users, "A"), login(app, users, "B")
    a.post("/api/profile", json=PROFILE)
    with redirect_stdout(io.StringIO()):
        sched.run_one(users.find("A").id)
    rows = a.get("/api/results").json()["rows"]
    assert all(r["query"] for r in rows)  # each job says which search found it
    gone = rows[0]["id"]

    assert a.post("/api/dismiss", json={"ids": []}).status_code == 400
    assert a.post("/api/dismiss", json={"ids": [gone]}).json() == {"dismissed": 1}
    after = a.get("/api/results").json()
    assert after["total"] == 2 and gone not in {r["id"] for r in after["rows"]}
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(a.get("/api/export.xlsx").content))
    head = [c.value for c in wb["All listings"][1]]
    assert gone not in {r[head.index("Job ID")].value for r in wb["All listings"].iter_rows(min_row=2)}
    assert b.post("/api/dismiss", json={"ids": [gone]}).json() == {"dismissed": 1}  # B's own list
    assert a.get("/api/results").json()["total"] == 2

    a.post("/api/dismiss", json={"ids": [gone], "undo": True})
    assert a.get("/api/results").json()["total"] == 3


def test_unscored_jobs_say_when_they_get_a_fit_or_why_not(svc, monkeypatch):
    from jev_job_hunter import scheduler as sch
    monkeypatch.setattr(sch, "MAX_OPEN", 1)
    app, users, sched, _, _ = svc
    a = login(app, users, "A")
    a.post("/api/profile", json=PROFILE)
    a.post("/api/watch", json={"enabled": True})
    with redirect_stdout(io.StringIO()):
        sched.run_one(users.find("A").id)  # search only: 2 jobs queued, 1 judged not worth reading
    rows = {r["id"]: r for r in a.get("/api/results").json()["rows"]}
    nxt = datetime.fromisoformat(sched.schedule(users.find("A").id)["next_read"])
    queued = sorted((r for r in rows.values() if r["eval"]["state"] == "queued"), key=lambda r: r["eval"]["at"])
    assert [datetime.fromisoformat(r["eval"]["at"]) - nxt for r in queued] == [timedelta(0), timedelta(minutes=12)]
    assert [r["eval"]["state"] for r in rows.values()].count("skipped") == 1

    a.post("/api/watch", json={"enabled": False})
    assert {r["eval"]["state"] for r in a.get("/api/results").json()["rows"]} == {"paused", "skipped"}
