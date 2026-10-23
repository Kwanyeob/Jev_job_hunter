"""Resume extraction (fake Claude client) and the local web server end to end (fake LinkedIn)."""

from __future__ import annotations

import json
import threading
import urllib.request
from http.server import ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from jev_job_hunter import resume, web
from jev_job_hunter.linkedin import run_linkedin

from test_linkedin import fake_client

PROFILE = {
    "candidate": {"headline": "Backend Engineer, 4 years Java", "summary": "Builds APIs.",
                  "years_experience": 4, "seniority": "mid", "skills": ["Java", "Python", "AWS", "PostgreSQL"],
                  "target_roles": ["Backend Engineer", "AI Engineer"], "interests": ["APIs"],
                  "locations": ["Seoul", "Remote"], "languages": ["English"], "education": ""},
    "search": {"queries": ["Backend Engineer"], "location": "Seoul", "remote": "any", "posted": "week"},
}


class FakeClient:
    def __init__(self, stop_reason="end_turn", text=json.dumps(PROFILE)):
        self.kw = None
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self._create))
        self._resp = SimpleNamespace(stop_reason=stop_reason,
                                     content=[SimpleNamespace(type="thinking"), SimpleNamespace(type="text", text=text)])

    def _create(self, **kw):
        self.kw = kw
        return self._resp


def test_extract_profile_request_shape():
    client = FakeClient()
    got = resume.extract_profile(b"%PDF-1.7 fake", client=client)
    assert got == PROFILE
    kw = client.kw
    assert kw["model"] == "claude-sonnet-5-5" and kw["output_config"]["effort"] == "medium"
    assert kw["fallbacks"] == "default" and kw["betas"] == ["server-side-fallback-2026-07-01"]
    assert "thinking" not in kw
    schema = kw["output_config"]["format"]["schema"]
    fields = set(schema["properties"]["candidate"]["properties"])
    assert not fields & {"name", "email", "phone", "address"}
    doc, text = kw["messages"][0]["content"]
    assert doc["type"] == "document" and doc["source"]["media_type"] == "application/pdf"
    assert "Today is" in text["text"]


def test_extract_profile_career_goal():
    client = FakeClient()
    resume.extract_profile(b"%PDF-1.7", client=client)
    plain = client.kw["messages"][0]["content"][1]["text"]
    assert "Career goal from the candidate" not in plain and "adjacent_roles" in plain
    fields = client.kw["output_config"]["format"]["schema"]["properties"]["candidate"]
    assert {"career_goal", "transferable_skills", "adjacent_roles"} <= set(fields["required"])

    resume.extract_profile(b"%PDF-1.7", goal="  PM/PO로\n전환 " + "x" * 500, client=client)
    text = client.kw["messages"][0]["content"][1]["text"]
    assert "Career goal from the candidate: PM/PO로 전환 x" in text and "x" * 301 not in text


def test_retarget_profile_sends_no_pdf():
    client = FakeClient()
    got = resume.retarget_profile({**PROFILE, "path": "/data/users/x/profile.yaml"}, "Product Owner", client=client)
    assert got == PROFILE
    (only,) = client.kw["messages"][0]["content"]
    assert only["type"] == "text"
    assert "Career goal from the candidate: Product Owner" in only["text"]
    assert '"Backend Engineer, 4 years Java"' in only["text"] and "/data/users" not in only["text"]


def test_mock_fit_counts_transferable_skills():
    from jev_job_hunter.jev import mock_fit
    job = {"title": "Technical Product Manager", "location": "Seoul",
           "text": "Own the roadmap, write specs, run A/B experiments with engineering and stakeholders."}
    dev = {**PROFILE["candidate"], "target_roles": ["Technical Product Manager"]}
    pivot = {**dev, "transferable_skills": ["Wrote specs with the PM team", "Ran A/B experiments",
                                            "Planned the roadmap with stakeholders"]}
    before = mock_fit({"candidate_profile": dev, "job": job})["skills_match"]
    after = mock_fit({"candidate_profile": pivot, "job": job})["skills_match"]
    assert after > before + 0.3


def test_extract_profile_errors():
    with pytest.raises(resume.ResumeError, match="not a PDF"):
        resume.extract_profile(b"hello", client=FakeClient())
    with pytest.raises(resume.ResumeError, match="declined"):
        resume.extract_profile(b"%PDF-1.7", client=FakeClient(stop_reason="refusal"))
    with pytest.raises(resume.ResumeError, match="cut off"):
        resume.extract_profile(b"%PDF-1.7", client=FakeClient(stop_reason="max_tokens"))


@pytest.fixture
def server(tmp_path, monkeypatch):
    (tmp_path / "config").mkdir()
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test")
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    monkeypatch.setattr(resume, "extract_profile", lambda pdf, client=None: json.loads(json.dumps(PROFILE)))

    def runner(root, **kw):
        return run_linkedin(root, delay=0, client=fake_client(), out_dir=root / "results", **kw)

    hunt = web.Hunt()
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), web.make_handler(tmp_path, hunt, runner=runner))
    httpd.daemon_threads = True
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{httpd.server_address[1]}", tmp_path, hunt
    httpd.shutdown()


def _post(url, body: bytes, ctype: str):
    req = urllib.request.Request(url, data=body, method="POST", headers={"Content-Type": ctype})
    with urllib.request.urlopen(req, timeout=10) as r:
        return json.loads(r.read())


def test_web_flow(server):
    base, root, hunt = server
    with urllib.request.urlopen(base + "/", timeout=5) as r:
        assert b"Drop your resume PDF" in r.read()
    state = json.loads(urllib.request.urlopen(base + "/api/state", timeout=5).read())
    assert state["claude"] is True and state["jev"] is False and state["profile"] is None

    got = _post(base + "/api/resume", b"%PDF-1.7 fake", "application/pdf")
    assert got["profile"]["candidate"]["target_roles"] == ["Backend Engineer", "AI Engineer"]
    assert (root / "config/profile.yaml").is_file()

    started = _post(base + "/api/hunt", json.dumps({"profile": got["profile"], "max_open": 5}).encode(),
                    "application/json")
    assert started["mock"] is True  # no Jev key → forced mock

    kinds, data = [], {}
    with urllib.request.urlopen(base + "/api/events?since=0", timeout=10) as r:
        kind = ""
        for raw in r:
            line = raw.decode("utf-8").rstrip("\n")
            if line.startswith("event: "):
                kind = line[7:]
            elif line.startswith("data: "):
                kinds.append(kind)
                data.setdefault(kind, []).append(json.loads(line[6:]))
                if kind == "done":
                    break
    assert kinds.count("card") == 3
    assert {"id", "title", "company", "href", "posted", "posted_text"} <= set(data["card"][0])
    assert "triage" in kinds and kinds.count("fit") == 2
    assert Path(data["done"][0]["xlsx"]).is_file()

    with urllib.request.urlopen(base + "/api/export.xlsx", timeout=10) as r:
        assert "attachment" in r.headers["Content-Disposition"]
        xlsx = r.read()
    out = root / "export.xlsx"
    out.write_bytes(xlsx)
    wb = load_workbook(out)
    assert len(list(wb["Matches"].iter_rows(min_row=2))) == 2
    assert len(list(wb["All listings"].iter_rows(min_row=2))) == 3


def test_web_rejects_empty_profile(server):
    base, _, _ = server
    req = urllib.request.Request(base + "/api/hunt", data=json.dumps({"profile": {}}).encode(), method="POST",
                                 headers={"Content-Type": "application/json"})
    with pytest.raises(urllib.error.HTTPError) as e:
        urllib.request.urlopen(req, timeout=5)
    assert e.value.code == 400
