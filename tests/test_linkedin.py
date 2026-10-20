"""LinkedIn flow over a fake HTTP transport: guest search → Jev triage → posting / ATS → Jev fit → xlsx."""

from __future__ import annotations

import io
import json
from contextlib import redirect_stdout
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

import httpx
import pytest
from openpyxl import load_workbook

from jev_job_hunter import linkedin
from jev_job_hunter.linkedin import (
    HEADERS, GuestClient, RateLimited, apply_target, is_authwall, parse_cards, parse_posting, run_linkedin,
    search_url,
)
from jev_job_hunter.profile import load_profile

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "fixtures"
JD = (
    "We are hiring a Backend Engineer to build Python and Go services on AWS with PostgreSQL "
    "and Kafka. 3+ years of experience. Seoul or remote. " * 8
)
CARDS = [
    {"id": "4001", "href": "https://www.linkedin.com/jobs/view/backend-engineer-4001",
     "title": "Backend Engineer", "company": "Toss", "location": "Seoul, South Korea", "posted": "2026-09-25", "posted_text": "3 days ago"},
    {"id": "4002", "href": "https://www.linkedin.com/jobs/view/ai-engineer-4002",
     "title": "AI Engineer, LLM Platform", "company": "Upstage", "location": "Seoul", "posted": "2026-09-24", "posted_text": "4 days ago"},
    {"id": "4003", "href": "https://www.linkedin.com/jobs/view/sales-4003",
     "title": "Account Executive", "company": "Acme", "location": "Seoul", "posted": "2026-09-24", "posted_text": "4 days ago"},
]
ATS_URL = "https://boards.greenhouse.io/toss/jobs/1"
APPLY_RAW = "https://www.linkedin.com/jobs/view/externalApply/4001?url=https%3A%2F%2Fboards.greenhouse.io%2Ftoss%2Fjobs%2F1&amp;urlHash=x"
AUTHWALL = "https://www.linkedin.com/authwall?trk=x"


def card_html(c: dict) -> str:
    return (f'<li><div class="base-card" data-entity-urn="urn:li:jobPosting:{c["id"]}">'
            f'<a class="base-card__full-link" href="{c["href"]}?position=1&amp;pageNum=0"></a>'
            f'<h3 class="base-search-card__title">\n {c["title"]}\n</h3>'
            f'<h4 class="base-search-card__subtitle"><a href="#"> {c["company"]} </a></h4>'
            f'<span class="job-search-card__location"> {c["location"]} </span>'
            f'<time class="job-search-card__listdate" datetime="{c["posted"]}">\n {c["posted_text"]}\n</time>'
            + (f'<img class="artdeco-entity-image" data-delayed-url="{c["logo"]}">' if c.get("logo") else "")
            + '</div></li>')


def posting_html(text: str, apply: str = "") -> str:
    code = f'<code id="applyUrl" style="display: none"><!--"{apply}"--></code>' if apply else ""
    return (f'<button class="apply-button">Apply<icon data-svg-class-name="apply-button__offsite-apply-icon-svg"></icon></button>{code}'
            f'<section class="show-more-less-html"><div class="show-more-less-html__markup"><p>{text}</p></div>'
            f'<button>Show more</button></section>'
            '<ul><li class="description__job-criteria-item"><h3 class="description__job-criteria-subheader"> Seniority level </h3>'
            '<span class="description__job-criteria-text"> Mid-Senior level </span></li></ul>')


POSTINGS = {
    "4001": posting_html("LinkedIn copy of the description.", APPLY_RAW),  # short → ATS
    "4002": posting_html(JD),                                             # Easy Apply only → LinkedIn text
    "4003": posting_html(JD),
}


class FakeLinkedIn:
    """httpx transport: guest search, guest postings, public /jobs/view/ pages, one ATS page."""

    def __init__(self, wall_on: str = "", rate_limit: int = 0, cards: list[dict] | None = None):
        self.visited: list[str] = []
        self.cards = CARDS if cards is None else cards
        self.wall_on = wall_on      # job id whose posting redirects to the authwall
        self.rate_limit = rate_limit  # first N requests answer 429

    def __call__(self, request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        self.visited.append(url)
        if self.rate_limit:
            self.rate_limit -= 1
            return httpx.Response(429)
        if url.startswith(AUTHWALL):
            return httpx.Response(200, text="<html>Sign in</html>")
        if url.startswith(linkedin.GUEST_SEARCH):
            start = int(parse_qs(urlsplit(url).query).get("start", ["0"])[0])
            return httpx.Response(200, text="".join(card_html(c) for c in self.cards[start:start + linkedin.PAGE_SIZE]))
        if url == ATS_URL:
            return httpx.Response(200, text=f"<html><body><h1>Backend Engineer</h1><p>{JD}</p></body></html>")
        for jid, page in POSTINGS.items():
            if url.endswith(f"/jobPosting/{jid}") or url.endswith(f"-{jid}"):
                if jid == self.wall_on:
                    return httpx.Response(302, headers={"Location": AUTHWALL})
                return httpx.Response(200, text=page)
        return httpx.Response(404)


def fake_client(site: FakeLinkedIn | None = None, delay: float = 0) -> GuestClient:
    return GuestClient(delay, http=httpx.Client(transport=httpx.MockTransport(site or FakeLinkedIn()),
                                                headers=HEADERS, follow_redirects=True))


@pytest.fixture
def profile(tmp_path):
    src = (ROOT / "config/profile.example.yaml").read_text(encoding="utf-8")
    p = tmp_path / "profile.yaml"
    p.write_text(src.replace("candidate:\n", "candidate:\n  name: Jane Doe\n  email: jane@example.com\n", 1),
                 encoding="utf-8")
    return p


def _hunt(tmp_path, profile, site: FakeLinkedIn) -> tuple[Path, str]:
    buf = io.StringIO()
    with redirect_stdout(buf):
        path = run_linkedin(ROOT, profile_path=str(profile), delay=0, mock=True,
                            out_dir=tmp_path / "out", client=fake_client(site))
    return path, buf.getvalue()


def test_profile_strips_pii_and_defaults(profile):
    prof = load_profile(ROOT, str(profile))
    assert "name" not in prof["candidate"] and "email" not in prof["candidate"]
    assert prof["search"]["queries"] == ["Backend Engineer", "AI Engineer"]
    assert prof["search"]["posted"] == "week"


def test_search_url_and_apply_target():
    url = search_url("AI Engineer", "South Korea", "week", "remote")
    assert url.startswith("https://www.linkedin.com/jobs/search/?keywords=AI+Engineer")
    assert "f_TPR=r604800" in url and "f_WT=2" in url and "location=South+Korea" in url
    assert "f_TPR" not in search_url("x", posted="any")
    assert apply_target(APPLY_RAW) == ATS_URL
    assert apply_target("https://jobs.lever.co/a/b") == "https://jobs.lever.co/a/b"
    assert apply_target("https://www.linkedin.com/jobs/view/4001") == ""
    assert apply_target("") == ""
    assert is_authwall("https://www.linkedin.com/authwall?trk=1")
    assert not is_authwall("https://www.linkedin.com/jobs/view/4001")


def test_parse_real_guest_search_page():
    cards = parse_cards((FIXTURES / "li_guest_search.html").read_text(encoding="utf-8"))
    assert len(cards) == 10
    first = cards[0]
    assert first["id"] == "4466644352"
    assert first["title"] == "Software Engineer, AI Model Enablement & Inference"
    assert first["company"] == "FuriosaAI" and first["location"] == "Seoul, Seoul, South Korea"
    assert first["posted"] == "2026-10-09" and first["posted_text"] == "3 days ago"
    assert first["href"].startswith("https://kr.linkedin.com/jobs/view/") and "?" not in first["href"]
    assert all(c["id"] and c["title"] and c["company"] for c in cards)
    assert first["logo"].startswith("https://media.licdn.com/") and "company-logo" in first["logo"]
    assert "&amp;" not in first["logo"]
    assert all(not c["logo"] or "company-logo" in c["logo"] for c in cards)  # never the grey placeholder


def test_parse_real_guest_posting_page():
    post = parse_posting((FIXTURES / "li_guest_posting.html").read_text(encoding="utf-8"))
    assert post["text"].startswith("About FuriosaAI")
    assert len(post["text"]) > 1000 and "Show more" not in post["text"] and "<" not in post["text"]
    assert post["criteria"]["Seniority level"] == "Not Applicable"
    assert post["criteria"]["Employment type"] == "Full-time"
    assert post["apply_type"] == "Company site" and post["apply"] == ""  # logged out: no apply URL


def test_linkedin_flow_writes_excel(tmp_path, profile):
    site = FakeLinkedIn()
    path, out = _hunt(tmp_path, profile, site)
    assert out.strip().endswith("ACTION END")
    assert "STAGE triage" in out and "Overall fit" in out
    # both queries run (3 cards < one page → one request each), cards deduped, sales card not opened, ATS followed
    assert sum(1 for u in site.visited if u.startswith(linkedin.GUEST_SEARCH)) == 2
    assert ATS_URL in site.visited
    assert not any(u.endswith("/jobPosting/4003") for u in site.visited)

    wb = load_workbook(path)
    assert wb.sheetnames == ["Matches", "All listings", "Profile"]
    rows = list(wb["Matches"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    fits = [r[1] for r in rows]
    assert fits == sorted(fits, reverse=True)
    by_title = {r[2]: r for r in rows}
    head = [c.value for c in wb["Matches"][1]]
    assert by_title["Backend Engineer"][head.index("Read from")] == "ATS"
    assert by_title["AI Engineer, LLM Platform"][head.index("Read from")] == "LinkedIn"
    ats_row = 2 + [r[2] for r in rows].index("Backend Engineer")
    apply_col = head.index("Apply / ATS") + 1
    assert wb["Matches"].cell(row=ats_row, column=apply_col).hyperlink.target == ATS_URL
    assert by_title["Backend Engineer"][head.index("Job ID")] == "4001"
    assert len(list(wb["All listings"].iter_rows(min_row=2))) == 3
    prof_rows = {r[0]: r[1] for r in wb["Profile"].iter_rows(values_only=True)}
    assert "candidate.email" not in prof_rows and "candidate.name" not in prof_rows

    data = json.loads(path.with_suffix(".json").read_text(encoding="utf-8"))
    assert len(data["matches"]) == 2


def test_authwall_skips_one_posting_and_still_writes(tmp_path, profile, monkeypatch):
    monkeypatch.setattr(linkedin, "LINKEDIN_OPEN_THRESHOLD", 0.0)
    path, out = _hunt(tmp_path, profile, FakeLinkedIn(wall_on="4003"))  # lowest-ranked card hits the wall
    assert "login wall" in out
    assert path is not None and path.is_file()
    rows = list(load_workbook(path)["Matches"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2


def test_search_pages_ten_at_a_time():
    many = [{**CARDS[0], "id": str(5000 + i), "title": f"Backend Engineer {i}",
             "href": f"https://www.linkedin.com/jobs/view/x-{5000 + i}"} for i in range(25)]
    site = FakeLinkedIn(cards=many)
    with redirect_stdout(io.StringIO()):
        got = linkedin.collect_cards(fake_client(site), "AI", "", "week", "any", limit=30)
    starts = [parse_qs(urlsplit(u).query)["start"][0] for u in site.visited]
    assert starts == ["0", "10", "20"]  # 10, 10, then 5 < a full page → stop
    assert len(got) == 25 and len({c["id"] for c in got}) == 25


def test_rate_limit_backs_off_then_recovers(monkeypatch):
    waits = []
    monkeypatch.setattr(linkedin, "RATE_LIMIT_BACKOFF", 60.0)
    client = fake_client(FakeLinkedIn(rate_limit=2))
    monkeypatch.setattr(client, "sleep", lambda s: s > 0 and waits.append(s))  # pacing at delay=0 asks for <= 0
    assert "data-entity-urn" in client.get(linkedin.GUEST_SEARCH + "?keywords=x&start=0")
    assert len(waits) == 2 and 48 <= waits[0] <= 78 and 96 <= waits[1] <= 156  # 60 s then 120 s, jittered

    client = fake_client(FakeLinkedIn(rate_limit=99))
    monkeypatch.setattr(client, "sleep", lambda s: None)
    with pytest.raises(RateLimited):
        client.get(linkedin.GUEST_SEARCH + "?keywords=x&start=0")


def test_persistent_rate_limit_stops_hunt(tmp_path, profile, monkeypatch):
    monkeypatch.setattr(linkedin, "RATE_LIMIT_BACKOFF", 0.0)
    path, out = _hunt(tmp_path, profile, FakeLinkedIn(rate_limit=99))
    assert path is None
    assert "rate limit" in out and out.strip().endswith("ACTION END")


def test_requests_are_spaced_irregularly(monkeypatch):
    gaps = []
    client = fake_client(delay=10.0)
    monkeypatch.setattr(client, "sleep", lambda s: gaps.append(s))
    for _ in range(40):
        client.get(linkedin.GUEST_SEARCH + "?keywords=x&start=0")
    assert len(gaps) == 39  # none before the first request
    assert all(5.0 <= g <= 46.0 for g in gaps)  # 0.6–1.6 × delay, plus an occasional 1–3 × delay linger
    assert len(set(gaps)) == len(gaps) and max(gaps) - min(gaps) > 5.0  # never a fixed cadence


def test_target_seniority_becomes_linkedins_experience_filter():
    from urllib.parse import parse_qs, urlsplit
    q = parse_qs(urlsplit(search_url("Java", "Toronto", "week", "any", ["mid_senior", "associate"])).query)
    assert q["f_E"] == ["3,4"]
    assert parse_qs(urlsplit(search_url("Java", experience=["director"])).query)["f_E"] == ["5,6"]
    assert "f_E" not in search_url("Java") and "f_E" not in search_url("Java", experience=["bogus"])


def test_profile_keeps_only_known_target_seniority():
    from jev_job_hunter.profile import normalize_profile
    prof = normalize_profile({"candidate": {"target_roles": ["Dev"],
                                            "target_seniority": ["mid_senior", "ceo", "associate"]}}, "t")
    assert prof["candidate"]["target_seniority"] == ["associate", "mid_senior"]
    empty = normalize_profile({"candidate": {"target_roles": ["Dev"], "target_seniority": ["ceo"]}}, "t")
    assert "target_seniority" not in empty["candidate"]
