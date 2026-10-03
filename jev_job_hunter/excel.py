"""LinkedIn hunt → .xlsx: Matches (Jev fit), All listings (Jev triage), Profile."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from jev_job_hunter.questions import FIT_SAVE_THRESHOLD

HEAD_FILL = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(bold=True, color="FFFFFF")
SAVE_FILL = PatternFill("solid", fgColor="DCFCE7")
LINK_FONT = Font(color="2563EB", underline="single")

MATCH_COLS = (
    ("Rank", 6, None), ("Overall fit", 11, "fit"), ("Title", 42, "title"), ("Company", 22, "company"),
    ("Location", 22, "location"), ("Role", 8, "role_match"), ("Skills", 8, "skills_match"),
    ("Seniority", 10, "seniority_match"), ("Location fit", 11, "location_match"),
    ("Level", 16, "level"), ("Apply via", 14, "apply_type"), ("Read from", 10, "source"),
    ("Posted", 12, "posted"), ("Updated", 14, "posted_text"),
    ("Apply / ATS", 16, "apply_url"), ("LinkedIn", 12, "href"), ("Job ID", 14, "id"), ("Query", 20, "query"),
)
LIST_COLS = (
    ("Open score", 11, "open"), ("Opened", 8, "opened"), ("Title", 42, "title"),
    ("Company", 22, "company"), ("Location", 22, "location"), ("Posted", 12, "posted"),
    ("Updated", 14, "posted_text"), ("LinkedIn", 12, "href"), ("Job ID", 14, "id"), ("Query", 20, "query"),
)
PCT_KEYS = {"fit", "role_match", "skills_match", "seniority_match", "location_match", "open"}
LINK_KEYS = {"apply_url", "href"}


def _header(ws, cols) -> None:
    for c, (name, width, _) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"


def _row(ws, r: int, cols, rec: dict, rank: int | None = None) -> None:
    for c, (_, _, key) in enumerate(cols, 1):
        val = rank if key is None else rec.get(key)
        cell = ws.cell(row=r, column=c)
        if key in LINK_KEYS:
            if val:
                cell.value, cell.hyperlink, cell.font = "open", val, LINK_FONT
            continue
        if key == "opened":
            val = "Y" if val else ""
        cell.value = val
        if key in PCT_KEYS and isinstance(val, (int, float)):
            cell.number_format = "0%"


def _scale(ws, col: int, last: int) -> None:
    if last < 2:
        return
    ref = f"{get_column_letter(col)}2:{get_column_letter(col)}{last}"
    ws.conditional_formatting.add(ref, ColorScaleRule(
        start_type="num", start_value=0, start_color="FCA5A5",
        mid_type="num", mid_value=0.5, mid_color="FDE68A",
        end_type="num", end_value=1, end_color="86EFAC"))


def write_xlsx(path: Path, matches: list[dict], listings: list[dict], meta: dict) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Matches"
    _header(ws, MATCH_COLS)
    for i, m in enumerate(matches, 1):
        _row(ws, i + 1, MATCH_COLS, m, rank=i)
        if (m.get("fit") or 0) >= FIT_SAVE_THRESHOLD:
            for c in range(1, 6):
                ws.cell(row=i + 1, column=c).fill = SAVE_FILL
    _scale(ws, 2, len(matches) + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(MATCH_COLS))}{max(1, len(matches) + 1)}"

    wl = wb.create_sheet("All listings")
    _header(wl, LIST_COLS)
    for i, rec in enumerate(listings, 2):
        _row(wl, i, LIST_COLS, rec)
    _scale(wl, 1, len(listings) + 1)
    wl.auto_filter.ref = f"A1:{get_column_letter(len(LIST_COLS))}{max(1, len(listings) + 1)}"

    wp = wb.create_sheet("Profile")
    wp.column_dimensions["A"].width, wp.column_dimensions["B"].width = 20, 90
    for r, (k, v) in enumerate(meta.items(), 1):
        wp.cell(row=r, column=1, value=k).font = Font(bold=True)
        wp.cell(row=r, column=2, value=", ".join(map(str, v)) if isinstance(v, list) else v)

    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:  # the previous file is open in Excel
        path = path.with_name(path.stem + "-new" + path.suffix)
        wb.save(path)
    return path
